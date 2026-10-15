#!/usr/bin/env python3
"""
One-time scraper: Michigan court list -> court codes.

The court-selection page is an Angular app. Rows have no <a href>; clicking a row
routes to /case-search/court/<CODE>. So for every row we click it, read the code
from the URL, and go back to the list.

It does NOT touch the case-search/CAPTCHA part of the site, so it is just ~230 clicks.

Run (same venv, Playwright + Chromium already installed):
    python get_court_codes.py                # full run, visible browser
    python get_court_codes.py --limit 3      # quick test: first 3 courts only
    python get_court_codes.py --headless     # no window

Output (all saved in the folder you run it from):
    court_codes.xlsx   <- the Excel sheet: County, Court Name, Description, 7-Year Filter, Code, URL
    court_codes.csv    working file used for resuming
    court_codes.json   same data
    court_api_dump/    any JSON the page fetched while loading the list (see note at the end)

It is resumable: rerun and it skips courts already in court_codes.csv.
"""

import argparse
import atexit
import csv
import json
import re
import sys
from pathlib import Path
from urllib.parse import unquote

from playwright.sync_api import sync_playwright

LIST_URL = "https://micourt.courts.michigan.gov/case-search/court-selection"
CODE_RE = re.compile(r"/case-search/court/([^/?#]+)")
ROWS = "app-court-selection-item tr"
FIRST_NAME_JS = """
() => {
  const e = document.querySelector('app-court-selection-item td.court-name');
  return e ? e.innerText.trim() : '';
}
"""
READ_ROWS_JS = """
() => [...document.querySelectorAll('app-court-selection-item tr')].map(tr => {
  const t = s => (tr.querySelector(s)?.innerText || '').replace(/\\s+/g, ' ').trim();
  return {
    county: t('.court-county'),
    court_name: t('.court-name'),
    description: t('.court-description'),
    has_7yr_filter: !!tr.querySelector('.court-set-aside i.bi-check-lg'),
  };
})
"""
FIELDS = ["county", "court_name", "description", "has_7yr_filter", "code", "url"]
OUT_CSV = Path("court_codes.csv")
OUT_JSON = Path("court_codes.json")
OUT_XLSX = Path("court_codes.xlsx")
API_DIR = Path("court_api_dump")


def key_of(r):
    return (r["county"], r["court_name"], r["description"])


def load_done():
    done = {}
    if OUT_CSV.exists():
        with OUT_CSV.open(newline="", encoding="utf-8") as f:
            for r in csv.DictReader(f):
                done[key_of(r)] = r
    return done


def append_row(row):
    new = not OUT_CSV.exists()
    with OUT_CSV.open("a", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS)
        if new:
            w.writeheader()
        w.writerow(row)


def write_xlsx(rows):
    """Write all collected courts to court_codes.xlsx (also runs if you press Ctrl+C)."""
    if not rows:
        return
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Courts"
    ws.append(["County", "Court Name", "Description", "Has 7-Year Filter", "Code", "URL"])
    for r in sorted(rows, key=lambda r: (r["county"], r["court_name"])):
        has = "Yes" if str(r["has_7yr_filter"]).lower() == "true" else "No"
        ws.append([r["county"], r["court_name"], r["description"], has, r["code"], r["url"]])
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, w in enumerate([22, 38, 55, 18, 14, 60], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    try:
        wb.save(OUT_XLSX)
        print(f"Excel saved: {OUT_XLSX.resolve()}  ({len(rows)} courts)")
    except PermissionError:
        alt = OUT_XLSX.with_name("court_codes_new.xlsx")
        wb.save(alt)
        print(f"{OUT_XLSX} is open in Excel - saved to {alt.resolve()} instead")


def open_list(page):
    page.goto(LIST_URL, wait_until="domcontentloaded")
    page.wait_for_selector(ROWS, timeout=30000)
    page.wait_for_function("() => { const e=document.querySelector('app-court-selection-item td.court-name');"
                           " return e && e.innerText.trim().length > 0; }", timeout=30000)


def wait_rows_changed(page, old_first):
    page.wait_for_function(
        "old => { const e=document.querySelector('app-court-selection-item td.court-name');"
        " const t=e ? e.innerText.trim() : ''; return t && t !== old; }",
        arg=old_first, timeout=20000)


def pager(page):
    return page.locator("ngb-pagination").first


def next_disabled(page):
    return pager(page).locator("li.page-item.disabled a[aria-label='Next']").count() > 0


def click_next(page):
    old = page.evaluate(FIRST_NAME_JS)
    pager(page).locator("a[aria-label='Next']").click()
    wait_rows_changed(page, old)


def go_to_page(page, n):
    """Walk from page 1 to page n (clicks the number if visible, otherwise 'Next')."""
    cur = 1
    while cur < n:
        link = pager(page).locator("a.page-link", has_text=re.compile(rf"^\s*{n}\s*$"))
        if link.count():
            old = page.evaluate(FIRST_NAME_JS)
            link.first.click()
            wait_rows_changed(page, old)
            return
        click_next(page)
        cur += 1


def ensure_list(page, pg):
    open_list(page)
    go_to_page(page, pg)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--headless", action="store_true")
    ap.add_argument("--limit", type=int, default=0, help="stop after N new courts (testing)")
    args = ap.parse_args()

    done = load_done()
    print(f"Already have {len(done)} courts in {OUT_CSV}")
    atexit.register(lambda: write_xlsx(list(done.values())))
    new_count = 0
    failures = []
    sniffing = {"on": True, "n": 0}

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=args.headless)
        ctx = browser.new_context(viewport={"width": 1300, "height": 900})
        page = ctx.new_page()

        def sniff(resp):
            if not sniffing["on"]:
                return
            try:
                if resp.request.resource_type in ("xhr", "fetch") and "json" in (resp.headers.get("content-type") or ""):
                    API_DIR.mkdir(exist_ok=True)
                    sniffing["n"] += 1
                    (API_DIR / f"{sniffing['n']:02d}.json").write_text(
                        json.dumps({"url": resp.url, "body": resp.json()}, indent=1), encoding="utf-8")
                    print(f"  [api] {resp.url}")
            except Exception:  # noqa: BLE001
                pass

        page.on("response", sniff)
        open_list(page)
        sniffing["on"] = False

        pg = 1
        stop = False
        while not stop:
            rows = page.evaluate(READ_ROWS_JS)
            print(f"--- list page {pg}: {len(rows)} rows ---")
            on_list = True
            for i, row in enumerate(rows):
                if key_of(row) in done:
                    continue
                ok = False
                for attempt in range(1, 4):
                    try:
                        if not on_list:
                            ensure_list(page, pg)
                            on_list = True
                        page.locator(ROWS).nth(i).click()
                        page.wait_for_url(CODE_RE, timeout=15000)
                        on_list = False
                        code = unquote(CODE_RE.search(page.url).group(1))
                        out = dict(row, code=code, url=page.url)
                        append_row(out)
                        done[key_of(row)] = out
                        new_count += 1
                        print(f"  p{pg} r{i + 1:02d}  {code:<10} {row['county']} | {row['court_name']}")
                        ok = True
                        break
                    except Exception as e:  # noqa: BLE001
                        print(f"  p{pg} r{i + 1:02d} attempt {attempt} failed: {str(e)[:100]}")
                        on_list = False
                if not ok:
                    failures.append(row)
                if args.limit and new_count >= args.limit:
                    stop = True
                    break
            if stop:
                break
            if not on_list:
                ensure_list(page, pg)
            if next_disabled(page):
                break
            click_next(page)
            pg += 1

        browser.close()

    all_rows = list(done.values())
    OUT_JSON.write_text(json.dumps(all_rows, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"\nDone. {len(all_rows)} courts total ({new_count} new this run) -> {OUT_CSV}, {OUT_JSON}, {OUT_XLSX}")
    if failures:
        print(f"{len(failures)} rows failed - just rerun the script, it will retry only those:")
        for r in failures:
            print("   ", r["county"], "|", r["court_name"])
    if sniffing["n"]:
        print(f"\nNote: the list page fetched {sniffing['n']} JSON response(s) -> {API_DIR}/. "
              f"If one of them contains the codes, send it to me and the scraper becomes a single request.")


if __name__ == "__main__":
    sys.exit(main())